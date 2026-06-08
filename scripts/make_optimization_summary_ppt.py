#!/usr/bin/env python3
"""Generate the HZVision Pose2Sim 阶段性优化总结 deck (.pptx).

Reproducible: regenerates the same slides every run. Content summarizes the
triangulation-stage stabilization work added on top of upstream Pose2Sim
(see docs/hzvision-integration.md and the project memory notes).

Usage:
    python scripts/make_optimization_summary_ppt.py [output.pptx]
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# --- palette ---------------------------------------------------------------
BLUE = RGBColor(0x1B, 0x5F, 0xA8)      # HZVision blue
DARK = RGBColor(0x20, 0x2A, 0x36)
GREY = RGBColor(0x5A, 0x64, 0x70)
LIGHT = RGBColor(0xEC, 0xF2, 0xF9)
GREEN = RGBColor(0x1E, 0x8E, 0x3E)
RED = RGBColor(0xC0, 0x39, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"   # CJK-friendly; viewers substitute if absent

SW, SH = Inches(13.333), Inches(7.5)   # 16:9

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def _set(run, size, color=DARK, bold=False, font=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def box(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def band(slide, color, l, t, w, h):
    sp = slide.shapes.add_shape(1, l, t, w, h)  # rectangle
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def header(slide, title, kicker=None):
    band(slide, BLUE, 0, 0, SW, Inches(1.15))
    band(slide, RGBColor(0x12, 0x44, 0x7A), 0, Inches(1.15), SW, Inches(0.06))
    _, tf = box(slide, Inches(0.55), Inches(0.18), Inches(12.2), Inches(0.85))
    p = tf.paragraphs[0]
    if kicker:
        r = p.add_run(); r.text = kicker + "  ·  "; _set(r, 16, RGBColor(0xCF, 0xE2, 0xF5), bold=False)
    r = p.add_run(); r.text = title; _set(r, 27, WHITE, bold=True)


def bullets(slide, items, left=Inches(0.7), top=Inches(1.5),
            width=Inches(12.0), height=Inches(5.6), size=18, gap=8):
    _, tf = box(slide, left, top, width, height)
    tf.word_wrap = True
    first = True
    for it in items:
        lvl = it[0] if isinstance(it, tuple) else 0
        text = it[1] if isinstance(it, tuple) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        p.level = lvl
        bullet = "▍ " if lvl == 0 else "•  "
        # color cues: support inline [g]...[/g] green, [r]...[/r] red
        prefix = p.add_run(); prefix.text = bullet
        _set(prefix, size, BLUE if lvl == 0 else GREY, bold=(lvl == 0))
        _emit_rich(p, text, size, bold=(lvl == 0))


def _emit_rich(p, text, size, bold=False):
    import re
    parts = re.split(r"(\[g\].*?\[/g\]|\[r\].*?\[/r\]|\[b\].*?\[/b\])", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("[g]"):
            r = p.add_run(); r.text = part[3:-4]; _set(r, size, GREEN, bold=True)
        elif part.startswith("[r]"):
            r = p.add_run(); r.text = part[3:-4]; _set(r, size, RED, bold=True)
        elif part.startswith("[b]"):
            r = p.add_run(); r.text = part[3:-4]; _set(r, size, DARK, bold=True)
        else:
            r = p.add_run(); r.text = part; _set(r, size, DARK if bold else RGBColor(0x33, 0x3B, 0x45), bold=bold)


# ===========================================================================
# 1. Title
# ===========================================================================
s = prs.slides.add_slide(BLANK)
band(s, BLUE, 0, 0, SW, SH)
band(s, RGBColor(0x12, 0x44, 0x7A), 0, Inches(5.0), SW, Inches(2.5))
_, tf = box(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(2.4))
p = tf.paragraphs[0]; r = p.add_run(); r.text = "Pose2Sim · HZVision 分支"; _set(r, 24, RGBColor(0xCF, 0xE2, 0xF5), bold=False)
p = tf.add_paragraph(); r = p.add_run(); r.text = "无标记动捕稳定性优化"; _set(r, 44, WHITE, bold=True)
p = tf.add_paragraph(); r = p.add_run(); r.text = "阶段性总结 —— 相比原始 Pose2Sim 的改进"; _set(r, 26, WHITE, bold=False)
_, tf = box(s, Inches(0.9), Inches(5.4), Inches(11.5), Inches(1.6))
for line in ["4 相机 PnVision / HZVision 生产管线", "弯腰 · 转身 · 行走 · 遮挡场景下的去抖动 / 去翻转 / 去瞬移",
             "全部改动默认关闭，不配置即与上游逐字节一致"]:
    p = tf.add_paragraph(); r = p.add_run(); r.text = "›  " + line; _set(r, 17, RGBColor(0xDD, 0xEA, 0xF7))

# ===========================================================================
# 2. 背景与痛点
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "背景：原始 Pose2Sim 在 4 相机生产数据上的痛点", "Why")
bullets(s, [
    "原始 Pose2Sim 逐 marker 独立三角化 + 大间隙 last-value 填充，对真实试验的几类系统性伪影没有护栏：",
    (1, "[r]腰部拉伸/塌缩[/r]：回归出的虚拟 Hip 点（HALPE_26 kp19）多视角不一致，弯腰时被拒成 NaN，last-value 填充造成 ~20 cm 整体瞬移"),
    (1, "[r]左右髋互换[/r]：行走时骨盆侧向轴沿相机深度被透视压缩 + 背向相机，几乎零代价就能 180° 翻转"),
    (1, "[r]朝向抖动/翻转[/r]：转身/侧站（foreshortening）时骨盆 yaw 病态，逐帧晃动甚至瞬时翻转"),
    (1, "[r]面部/头部跳变[/r]：低头时脸部 keypoint 置信度边缘化，参与三角化的相机集合逐帧切换"),
    (1, "[r]肢体爆炸[/r]：两相机分歧时手臂/肩膀独立三角化炸到 1–2 m"),
    (1, "[r]2D 抖动 / 整帧清空[/r]：静止时脚/髋逐帧抖；弯腰瞬时低分把整帧 2D 清空"),
    "目标：在不破坏上游默认行为的前提下，从[b]根因[/b]逐级修复，并用可复现脚本验证。",
], size=17, gap=7)

# ===========================================================================
# 3. 改进全景
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "改进全景：按管线阶段叠加的稳定能力", "Overview")
rows = [
    ("阶段", "新增能力", "解决的问题"),
    ("2D 姿态", "时序平滑(one_euro) + 不因均分低清空整帧 + 单人选择", "逐帧脚/髋抖动；弯腰遮挡整帧丢失"),
    ("入口/吞吐", "avi2trc CLI + 批量推理 + 并行三角化 + mosaic 存视频", "一条命令 AVI→TRC；提速；可视化质检"),
    ("三角化", "刚体 marker 组(pelvis/head) + Procrustes/Kabsch 模板", "骨盆拉伸/塌缩 + 间隙瞬移"),
    ("三角化", "手性护栏 + 朝向保持(hold)", "行走时左右髋互换"),
    ("三角化", "鲁棒 SO(3) 朝向平滑(测地中值)", "foreshortening 下 yaw 尖峰/翻转"),
    ("三角化", "解剖学扭转钳制", "骨盆与躯干无耦合的过度旋转"),
    ("三角化", "非物理肢体护栏", "手臂/肩膀炸到 1–2 m"),
]
left, top, w = Inches(0.55), Inches(1.45), Inches(12.25)
table = s.shapes.add_table(len(rows), 3, left, top, w, Inches(5.4)).table
table.columns[0].width = Inches(1.9)
table.columns[1].width = Inches(5.7)
table.columns[2].width = Inches(4.65)
for ci, name in enumerate(rows[0]):
    cell = table.cell(0, ci)
    cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = cell.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = name; _set(r, 15, WHITE, bold=True)
for ri in range(1, len(rows)):
    for ci in range(3):
        cell = table.cell(ri, ci)
        cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT if ri % 2 else WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_top = Pt(3); cell.margin_bottom = Pt(3)
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = rows[ri][ci]
        _set(r, 13.5, DARK if ci == 0 else RGBColor(0x33, 0x3B, 0x45), bold=(ci == 0))

# ===========================================================================
# 4. 2D 阶段
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "① 2D 姿态阶段：先稳住输入再三角化", "2D pose")
bullets(s, [
    "[b]时序去抖（one_euro）[/b]：在 RTMPose 推理、pose-NMS、person tracking 之后、写 JSON / 画检测视频之前执行，让 JSON、可视化、三角化共用同一套稳定 2D 点。",
    (1, "推荐 min_cutoff=1.0, beta=0.02, d_cutoff=1.0, max_gap=5"),
    "[b]不因整帧均分低而清空[/b]（drop_low_average_pose=false）：弯腰/遮挡/手脚出画时，单点低分不再清空整帧；各点原始置信度保留，后续按 likelihood 逐点过滤。",
    (1, "否则弯腰瞬时低分→整帧 2D 空→TRC 用 last-value 填→看起来头还停在上面"),
    "[b]单人选择 / 关联[/b]：多视角下稳定挑出同一个人，避免跨相机选错人。",
    "[b]站立片段 3D 稳定配置[/b]：收紧 reproj 阈值到 ~10px、最低相机数=3，避免单 marker 退化成低误差的两相机解。",
    "原则：默认关闭；不开启时输出不变。",
], size=17, gap=8)

# ===========================================================================
# 5. avi2trc / 吞吐
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "② 入口与吞吐：一条命令 AVI → TRC", "Pipeline")
bullets(s, [
    "[b]avi2trc CLI[/b]：给 HZVision 主程序调用的统一入口，不再手工替换虚拟环境里的 pose2sim 文件夹。",
    (1, "avi2trc --trial-dir /path/Session/Trial_1 --batch-size 16 --overwrite-pose"),
    (1, "trial 下 videos/*.avi，session 下 calibration/*.toml；可 Python 直接调用"),
    "[b]批量姿态推理[/b]：成批送入 RTMPose，提高 GPU 利用率与吞吐。",
    "[b]并行三角化[/b]：多帧/多机位并行，缩短整段处理时间。",
    "[b]mosaic 存视频[/b]：多机位检测结果拼接成一张对照视频，便于人工质检与回归对比。",
    "[b]分支维护[/b]：保留上游完整提交历史；upstream / hzvision 双远端，便于持续同步官方更新。",
], size=18, gap=9)

# ===========================================================================
# 6. 刚体组 + Procrustes
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "③ 刚体 marker 组 + Procrustes/Kabsch 模板", "Triangulation")
bullets(s, [
    "[b]核心思想[/b]：骨盆、颅骨在解剖上是刚体——用刚体拟合稳定其轨迹，把组内距离钉死，消除"
    "拉伸/塌缩与间隙瞬移。复用同一套刚体机制于 pelvis 与 head 两组。",
    (1, "pelvis=[Hip,RHip,LHip]，head=[Head,Nose,REye,LEye,REar,LEar]（颅骨刚体，排除会屈的 Neck）"),
    (1, "blend=1.0 + 关闭修正上限 → 纯刚体重建；reproj_error_threshold 作安全网，拟合解释不了的帧回退独立三角化"),
    "[b]逐组参数覆盖[/b]：每个刚体组可独立设阈值/平滑/方法，互不影响。",
    "[b]Procrustes / Kabsch 模板（关键修复）[/b]：旧模板对居中位置逐轴取中值，身体转动时不同朝向的 marker 互相抵消，把形状[r]压缩[/r]。",
    (1, "改为先用 Kabsch 把每个稳定帧旋到公共参考帧再取中值 → 无论是否转动都保住真实尺寸"),
    (1, "[g]yh644406 静态骨盆宽度 14.5 cm → 22.4 cm（恢复真实值）[/g]；已正面的试验变化 <0.6 cm"),
], size=16.5, gap=7)

# ===========================================================================
# 7. 手性护栏
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "④ 手性护栏 + 朝向保持：根治左右髋互换", "Triangulation")
bullets(s, [
    "[b]机理[/b]：骨盆三点近似共面，侧向轴沿相机深度被透视压缩 + 背向相机时，刚体拟合 180° yaw 翻转的 2D 代价几乎为零 → 自信地把 RHip/LHip 写反。",
    (1, "独立三角化是逐 marker 的（RHip 只用 RHip-2D，最坏 NaN，[b]永不换边[/b]）——因此它是可信的左右基准"),
    "[b]手性护栏 chirality_guard[/b]：刚体 (R−L) 向量与独立基准 (R−L) 反向的帧，拒绝刚体结果。",
    (1, "基准为 NaN 时退而用[b]时间连续性[/b]：相邻 30fps 帧 yaw 不可能突变 180°"),
    "[b]朝向保持 chirality_hold[/b]：拒绝帧不回退到退化的独立解（会塌缩骨盆宽度），而是把上一可信朝向按当前质心重定位后保持。",
    "[g]量化：行走试验 swap 帧 → 0[/g]（zjg_action02 16→0, wxs 18→0；交付集 mx 4→0, wxs2 8→0, zf 50→0, zjg 20→0）。",
    "非互换帧与改前逐字节一致；弯腰试验骨盆手性拒绝=0（不误伤）。",
], size=16.5, gap=7)

# ===========================================================================
# 8. 鲁棒朝向 + 扭转钳制
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "⑤ 鲁棒 SO(3) 朝向平滑 + 解剖学扭转钳制", "Triangulation")
bullets(s, [
    "[b]问题[/b]：侧站/转身（foreshortening）时骨盆 yaw 病态，残留两种伪影——接受帧的 yaw 尖峰/翻转，与拒绝帧回退独立解时的宽度塌缩。",
    "[b]鲁棒 SO(3) 平滑 smoothing_method=robust[/b]：平移仍 median_then_mean；旋转改在 SO(3) 上取窗口内[b]测地中值点(geodesic medoid)[/b]，容差内成员取弦平均。",
    (1, "逐分量中值无法剔除轴角向量上的非线性翻转；测地中值能投票淘汰尖峰、保留渐变真实运动"),
    "[b]解剖学扭转钳制 twist_guard[/b]：刚体骨盆与躯干无耦合会"
    "「像光滑的杆」过度旋转。以肩线为参考，绕竖直轴把骨盆相对躯干扭转钳制在 max_twist_deg。",
    (1, "真实躯干旋转约 45°，阈值取 50° → 只在伪影帧触发，干净试验上是 NO-OP"),
    "[g]量化：转身静态试验「宽度塌缩 + 左右翻转」→ 0（zf 30/3→0/0, zjg 7/5→0/0）[/g]；",
    (1, "[g]行走 swap 12→0 且运动不被过平滑（yaw 范围 −180..165° 不变，中位 |骨盆−肩| 1.3°）[/g]"),
], size=16, gap=6)

# ===========================================================================
# 9. 肢体护栏
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "⑥ 非物理肢体护栏", "Triangulation")
bullets(s, [
    "[b]问题[/b]：两相机分歧时，独立三角化的手臂/肩膀可能炸到 1–2 m（被两路一致但错误的相机拉飞，reproj 误差反而低）。",
    "[b]做法[/b]：沿骨骼树（root→leaf 一遍）检查每根骨段；同时满足"
    "「> max_ratio × 试验中值长度」且「绝对超出 min_excess_m」才判为非物理。",
    (1, "确定性地把[b]远端子节点[/b]沿当前骨向拉回中值长度（锚定已修复的父节点）——比 NaN+插值对整条肢体/持续爆炸更稳"),
    (1, "+0.15 m 绝对下限 → 脚部正常抖动（踝-跟 ~5.5cm）不被误伤；刚体组 marker 受保护不被移动"),
    "[b]踩过的坑[/b]：按 reproj 误差丢点是反的（坏点误差低）→ 总是修远端；纯 ratio 无绝对下限会误伤脚部。",
    "[g]量化：lhl 116 cm 前臂 → 60 cm（粗爆炸消除）；静态/快速/行走试验 ON-vs-OFF 逐点 0.000 mm（真 NO-OP）[/g]",
], size=16.5, gap=7)

# ===========================================================================
# 10. 量化结果汇总
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "量化结果汇总（真实 PnVision 试验，改前 → 改后）", "Results")
rows = [
    ("场景 / 试验", "指标", "改前 → 改后"),
    ("弯腰 身体前屈 (面部)", "Nose 最大跳变 / NaN 帧", "13.2→6.1 cm / 96→6"),
    ("弯腰 身体前屈 (骨盆)", "整盆单帧瞬移 (f266/f479)", "21.9→0.4 / 20.0→1.3 cm"),
    ("行走 action02", "左右髋互换帧 (zjg/wxs)", "16→0 / 18→0"),
    ("行走 交付集", "swap (mx/wxs2/zf/zjg)", "4/8/50/20 → 0"),
    ("转身静态", "宽度塌缩 / 翻转 (zf/zjg)", "30·3 / 7·5 → 0"),
    ("静态骨盆宽度 yh644406", "Procrustes 模板恢复", "14.5 → 22.4 cm"),
    ("肢体爆炸 lhl", "前臂长度", "116 → 60 cm"),
    ("干净 / 快速 / 行走试验", "护栏 ON-vs-OFF 逐点差", "0.000 mm (NO-OP)"),
]
left, top = Inches(1.1), Inches(1.5)
table = s.shapes.add_table(len(rows), 3, left, top, Inches(11.1), Inches(5.3)).table
table.columns[0].width = Inches(4.1)
table.columns[1].width = Inches(3.6)
table.columns[2].width = Inches(3.4)
for ci, name in enumerate(rows[0]):
    cell = table.cell(0, ci); cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = cell.text_frame.paragraphs[0].add_run(); r.text = name; _set(r, 15, WHITE, bold=True)
for ri in range(1, len(rows)):
    for ci in range(3):
        cell = table.cell(ri, ci)
        cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT if ri % 2 else WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
        r = cell.text_frame.paragraphs[0].add_run(); r.text = rows[ri][ci]
        _set(r, 14, GREEN if ci == 2 else RGBColor(0x33, 0x3B, 0x45), bold=(ci == 2))

# ===========================================================================
# 11. 工程原则
# ===========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "贯穿始终的工程原则", "Principles")
bullets(s, [
    "[b]默认关闭、向后兼容[/b]：所有新功能 opt-in；不配置刚体组 / 护栏时，输出与上游 Pose2Sim 逐字节一致。",
    "[b]从根因修，不在报告里粉饰[/b]：每个伪影都用真实代码路径复现，定位到底是 2D 检测层还是三角化层（弯腰面部=真遮挡，骨盆瞬移=三角化层）。",
    "[b]先复现、再修、给前后证据[/b]：同一输入跑改前/改后，保留 TRC、计数、diff 作为证据。",
    "[b]可复现验证脚本[/b]：verify_pelvis_jitter / verify_limb_guard / sweep_artifacts / batch_validate / render_skeleton / audit_2d。",
    "[b]快迭代环路[/b]：从缓存的 2D JSON 重三角化，不重跑姿态检测，分钟级验证一个配置。",
    "[b]度量的坑也要修[/b]：TRC marker 顺序 RShoulder=21/LShoulder=24，曾误用 22/25(肘) 虚高 swap 计数——代码按名字解析始终正确，只修离线指标。",
], size=17, gap=10)

# ===========================================================================
# 12. 结语
# ===========================================================================
s = prs.slides.add_slide(BLANK)
band(s, BLUE, 0, 0, SW, SH)
_, tf = box(s, Inches(0.9), Inches(1.4), Inches(11.5), Inches(1.6))
p = tf.paragraphs[0]; r = p.add_run(); r.text = "小结"; _set(r, 40, WHITE, bold=True)
bullets_white = [
    "在不动上游默认行为的前提下，为 4 相机生产管线补齐了一整套三角化稳定护栏：",
    "刚体组 + Procrustes 模板 · 手性护栏 + hold · 鲁棒 SO(3) 朝向 · 扭转钳制 · 非物理肢体护栏",
    "弯腰瞬移、左右髋互换、转身抖动/翻转、肢体爆炸——在真实试验上量化清零，运动不被过平滑。",
    "后续方向：≥3 背向相机时手性基准缺失的残留兜底；滤波层逐轴 Butterworth 破坏刚性的 ~3cm 宽度凹陷；持续同步官方更新。",
]
_, tf = box(s, Inches(0.9), Inches(2.9), Inches(11.6), Inches(3.6))
first = True
for t in bullets_white:
    p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
    p.space_after = Pt(14)
    r = p.add_run(); r.text = "›  "; _set(r, 18, RGBColor(0xCF, 0xE2, 0xF5), bold=True)
    r = p.add_run(); r.text = t; _set(r, 18, WHITE)
_, tf = box(s, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.7))
p = tf.paragraphs[0]; r = p.add_run()
r.text = "细节见 docs/hzvision-integration.md ·  分支 pose2sim-hzvision"
_set(r, 14, RGBColor(0xCF, 0xE2, 0xF5))

out = sys.argv[1] if len(sys.argv) > 1 else "Pose2Sim_HZVision_优化总结.pptx"
prs.save(out)
print(f"saved: {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
